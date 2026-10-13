"""Fixed JSON-to-DOCX worker. No user code, templates, file paths or URL fetching."""
import io
import json
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def export(payload):
    content=payload["content"]
    if len(content["title"]) > 120 or not 1 <= len(content["sections"]) <= 20:
        raise ValueError("invalid document bounds")
    document=Document()
    document.core_properties.author=""
    document.core_properties.last_modified_by=""
    document.core_properties.title=content["title"]
    for name in ("Normal","Title","Heading 1","Heading 2"):
        style=document.styles[name]
        style.font.name="Noto Sans CJK SC"
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC")
    document.styles["Normal"].font.size=Pt(11)
    document.add_paragraph(content["title"],style="Title")
    document.add_paragraph(f"版本 {int(payload['revision'])}　来源核查资料")
    for section in content["sections"]:
        if len(section["body"])>5000 or not 1 <= len(section["citations"]) <= 30:
            raise ValueError("invalid section bounds")
        document.add_heading(section["heading"],level=1)
        for line in section["body"].splitlines():
            document.add_paragraph(line)
        document.add_paragraph("来源与证据",style="Heading 2")
        for citation in section["citations"]:
            document.add_paragraph(citation["source_reference"])
            document.add_paragraph("证据标识："+citation["evidence_id"])
            if citation.get("quoted_text"):
                document.add_paragraph("原文摘录："+citation["quoted_text"])
    if content.get("user_notes"):
        document.add_heading("用户批注",level=1)
        document.add_paragraph("以下为用户个人批注，不属于检索验证的事实。")
        document.add_paragraph(content["user_notes"])
    buffer=io.BytesIO()
    document.save(buffer)
    result=buffer.getvalue()
    if len(result)>2_000_000:
        raise ValueError("export too large")
    return result


if __name__=="__main__":
    try:
        raw=sys.stdin.buffer.read(100_001)
        if len(raw)>100_000:
            raise ValueError("input too large")
        sys.stdout.buffer.write(export(json.loads(raw)))
    except Exception:
        print("document-export-worker-failed",file=sys.stderr)
        raise SystemExit(1)
