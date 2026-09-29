from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

from app.domain.documents.document_ingestion import (
    DocumentArtifact,
    DocumentFormat,
    DocumentFormatMismatchError,
    DocumentReadError,
    EvidenceLocation,
    InvalidSpreadsheetError,
    NoExtractableContentError,
    ParsedDocument,
    ParsedFragment,
)


class XlsxDocumentParser:
    document_format: DocumentFormat = "xlsx"

    def __init__(self, *, include_hidden_sheets: bool = False) -> None:
        self._include_hidden_sheets = include_hidden_sheets

    def parse(self, artifact: DocumentArtifact) -> ParsedDocument:
        if artifact.document_format != self.document_format:
            raise DocumentFormatMismatchError(
                "XlsxDocumentParser only accepts xlsx artifacts"
            )

        workbook = self._open_workbook(artifact.path)
        try:
            fragments = self._extract_fragments(workbook)
        finally:
            workbook.close()

        if not fragments:
            raise NoExtractableContentError(
                "XLSX workbook contains no extractable visible rows"
            )

        return ParsedDocument(
            artifact=artifact,
            fragments=tuple(fragments),
        )

    @staticmethod
    def _open_workbook(path: Path) -> Workbook:
        if not path.is_file():
            raise DocumentReadError(
                f"XLSX document could not be read: {path}"
            )
        try:
            return load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
            )
        except OSError as error:
            raise DocumentReadError(
                f"XLSX document could not be read: {path}"
            ) from error
        except (
            InvalidFileException,
            BadZipFile,
            KeyError,
            ValueError,
        ) as error:
            raise InvalidSpreadsheetError(
                "XLSX file is not a valid readable workbook"
            ) from error

    def _extract_fragments(
        self,
        workbook: Workbook,
    ) -> list[ParsedFragment]:
        fragments: list[ParsedFragment] = []

        for worksheet in workbook.worksheets:
            if (
                not self._include_hidden_sheets
                and worksheet.sheet_state != "visible"
            ):
                continue

            for row in worksheet.iter_rows():
                values = [
                    self._format_cell_value(cell)
                    for cell in row
                ]
                populated_indexes = [
                    index
                    for index, value in enumerate(values)
                    if value
                ]
                if not populated_indexes:
                    continue

                first_index = populated_indexes[0]
                last_index = populated_indexes[-1]
                row_text = " | ".join(
                    values[first_index : last_index + 1]
                )
                first_cell = row[first_index]
                last_cell = row[last_index]
                cell_reference = first_cell.coordinate
                if first_cell.coordinate != last_cell.coordinate:
                    cell_reference = (
                        f"{first_cell.coordinate}:{last_cell.coordinate}"
                    )

                fragments.append(
                    ParsedFragment(
                        ordinal=len(fragments),
                        text=row_text,
                        location=EvidenceLocation(
                            sheet_name=worksheet.title,
                            cell_reference=cell_reference,
                        ),
                    )
                )

        return fragments

    @staticmethod
    def _format_cell_value(
        cell: ReadOnlyCell | EmptyCell,
    ) -> str:
        value = cell.value
        if value is None:
            return ""
        if isinstance(value, datetime):
            if is_datetime(cell.number_format) == "date":
                return value.date().isoformat()
            return value.isoformat()
        if isinstance(value, (date, time)):
            return value.isoformat()
        return " ".join(str(value).split())
