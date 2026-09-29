from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from pptx import Presentation
from pptx.exc import InvalidXmlError, PackageNotFoundError
from pptx.presentation import Presentation as PptxPresentation

from app.domain.documents.document_ingestion import (
    DocumentArtifact,
    DocumentFormat,
    DocumentFormatMismatchError,
    DocumentReadError,
    EvidenceLocation,
    InvalidPresentationError,
    NoExtractableContentError,
    ParsedDocument,
    ParsedFragment,
)


class PptxDocumentParser:
    document_format: DocumentFormat = "pptx"

    def parse(self, artifact: DocumentArtifact) -> ParsedDocument:
        if artifact.document_format != self.document_format:
            raise DocumentFormatMismatchError(
                "PptxDocumentParser only accepts pptx artifacts"
            )

        presentation = self._open_presentation(artifact.path)
        fragments = self._extract_fragments(presentation)
        if not fragments:
            raise NoExtractableContentError(
                "PPTX presentation contains no extractable text"
            )

        return ParsedDocument(
            artifact=artifact,
            fragments=tuple(fragments),
        )

    @staticmethod
    def _open_presentation(path: Path) -> PptxPresentation:
        if not path.is_file():
            raise DocumentReadError(
                f"PPTX document could not be read: {path}"
            )
        try:
            return Presentation(path)
        except OSError as error:
            raise DocumentReadError(
                f"PPTX document could not be read: {path}"
            ) from error
        except (
            PackageNotFoundError,
            InvalidXmlError,
            BadZipFile,
            KeyError,
            ValueError,
        ) as error:
            raise InvalidPresentationError(
                "PPTX file is not a valid readable presentation"
            ) from error

    def _extract_fragments(
        self,
        presentation: PptxPresentation,
    ) -> list[ParsedFragment]:
        fragments: list[ParsedFragment] = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            title_shape = slide.shapes.title
            title_text = (
                self._normalize_text(title_shape.text)
                if title_shape is not None
                else ""
            )
            heading_path = (title_text,) if title_text else ()
            table_number = 0

            for shape in sorted(
                slide.shapes,
                key=lambda item: (item.top, item.left),
            ):
                if shape.has_table:
                    table_number += 1
                    for row_number, row in enumerate(
                        shape.table.rows,
                        start=1,
                    ):
                        values = [
                            self._normalize_text(cell.text)
                            for cell in row.cells
                        ]
                        populated_indexes = [
                            index
                            for index, value in enumerate(values)
                            if value
                        ]
                        if not populated_indexes:
                            continue
                        first_index = populated_indexes[0]
                        last_index = populated_indexes[-1]
                        self._append_fragment(
                            fragments,
                            " | ".join(
                                values[first_index : last_index + 1]
                            ),
                            slide_number=slide_number,
                            heading_path=heading_path,
                            table_number=table_number,
                            table_row_number=row_number,
                        )
                    continue

                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = self._normalize_text(paragraph.text)
                    if not text:
                        continue
                    self._append_fragment(
                        fragments,
                        text,
                        slide_number=slide_number,
                        heading_path=heading_path,
                    )

        return fragments

    @staticmethod
    def _normalize_text(text: str) -> str:
        return " ".join(text.split())

    @staticmethod
    def _append_fragment(
        fragments: list[ParsedFragment],
        text: str,
        *,
        slide_number: int,
        heading_path: tuple[str, ...],
        table_number: int | None = None,
        table_row_number: int | None = None,
    ) -> None:
        fragments.append(
            ParsedFragment(
                ordinal=len(fragments),
                text=text,
                location=EvidenceLocation(
                    heading_path=heading_path,
                    slide_number=slide_number,
                    table_number=table_number,
                    table_row_number=table_row_number,
                ),
            )
        )
