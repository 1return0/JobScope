import unittest
from dataclasses import replace
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from app.application.documents.document_parsing import (
    DocumentParserRegistry,
    DocumentParsingService,
)
from app.domain.documents.document_ingestion import (
    DocumentParsingError,
    inspect_document_artifact,
)
from app.infrastructure.documents.xlsx_document_parser import (
    XlsxDocumentParser,
)


class XlsxDocumentParserTest(unittest.TestCase):
    def test_extracts_rows_with_sheet_and_cell_range(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "positions.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Engineering Roles"
            sheet.append(["Role", "Degree", "Location"])
            sheet.append(["AI Agent Intern", None, "Shanghai"])
            sheet.append([None, None, None])
            sheet.append(["Deadline", date(2026, 10, 1)])
            workbook.save(path)
            artifact = inspect_document_artifact(
                path,
                "https://careers.example/positions.xlsx",
            )

            parsed = XlsxDocumentParser().parse(artifact)

        self.assertEqual(
            [
                "Role | Degree | Location",
                "AI Agent Intern |  | Shanghai",
                "Deadline | 2026-10-01",
            ],
            [fragment.text for fragment in parsed.fragments],
        )
        self.assertEqual(
            ["A1:C1", "A2:C2", "A4:B4"],
            [
                fragment.location.cell_reference
                for fragment in parsed.fragments
            ],
        )
        self.assertTrue(
            all(
                fragment.location.sheet_name == "Engineering Roles"
                for fragment in parsed.fragments
            )
        )

    def test_skips_hidden_sheets_by_default(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "positions.xlsx"
            workbook = Workbook()
            workbook.active.append(["Public role"])
            hidden = workbook.create_sheet("Internal Helper")
            hidden.sheet_state = "hidden"
            hidden.append(["Internal lookup value"])
            workbook.save(path)
            artifact = inspect_document_artifact(path, "source-xlsx")

            parsed = XlsxDocumentParser().parse(artifact)

        self.assertEqual(
            ["Public role"],
            [fragment.text for fragment in parsed.fragments],
        )

    def test_registry_dispatches_xlsx_parser(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "positions.xlsx"
            workbook = Workbook()
            workbook.active.append(["Product Intern", "Beijing"])
            workbook.save(path)
            artifact = inspect_document_artifact(path, "source-xlsx")

            parsed = DocumentParserRegistry(
                [XlsxDocumentParser()]
            ).parse(artifact)

        self.assertEqual(
            "Product Intern | Beijing",
            parsed.fragments[0].text,
        )

    def test_corrupt_xlsx_has_stable_failure_classification(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "corrupt.xlsx"
            path.write_bytes(b"not-an-office-workbook")
            artifact = inspect_document_artifact(path, "source-corrupt")
            service = DocumentParsingService(
                DocumentParserRegistry([XlsxDocumentParser()])
            )

            result = service.parse(artifact)

        self.assertEqual(
            "invalid-spreadsheet",
            result.failure.code,
        )

    def test_rejects_non_xlsx_artifact_when_called_directly(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "positions.xlsx"
            workbook = Workbook()
            workbook.active.append(["Content"])
            workbook.save(path)
            xlsx_artifact = inspect_document_artifact(
                path,
                "source-xlsx",
            )
            pptx_artifact = replace(
                xlsx_artifact,
                document_format="pptx",
            )

            with self.assertRaisesRegex(
                DocumentParsingError,
                "only accepts xlsx",
            ):
                XlsxDocumentParser().parse(pptx_artifact)


if __name__ == "__main__":
    unittest.main()
