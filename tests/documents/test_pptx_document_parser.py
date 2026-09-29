import unittest
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches

from app.application.documents.document_parsing import (
    DocumentParserRegistry,
    DocumentParsingService,
)
from app.domain.documents.document_ingestion import (
    DocumentParsingError,
    inspect_document_artifact,
)
from app.infrastructure.documents.pptx_document_parser import (
    PptxDocumentParser,
)


class PptxDocumentParserTest(unittest.TestCase):
    def test_extracts_slide_text_and_table_rows_with_locations(
        self,
    ) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "recruitment.pptx"
            presentation = Presentation()
            first_slide = presentation.slides.add_slide(
                presentation.slide_layouts[1]
            )
            first_slide.shapes.title.text = "Campus Recruitment"
            text_frame = first_slide.placeholders[1].text_frame
            text_frame.text = "Open to graduating students."
            text_frame.add_paragraph().text = "Apply before October."

            second_slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )
            second_slide.shapes.title.text = "Open Positions"
            table = second_slide.shapes.add_table(
                rows=2,
                cols=2,
                left=Inches(1),
                top=Inches(2),
                width=Inches(6),
                height=Inches(1.5),
            ).table
            table.cell(0, 0).text = "Role"
            table.cell(0, 1).text = "Location"
            table.cell(1, 0).text = "AI Agent Intern"
            table.cell(1, 1).text = "Shanghai"
            presentation.save(path)
            artifact = inspect_document_artifact(
                path,
                "https://careers.example/recruitment.pptx",
            )

            parsed = PptxDocumentParser().parse(artifact)

        self.assertEqual(
            [
                "Campus Recruitment",
                "Open to graduating students.",
                "Apply before October.",
                "Open Positions",
                "Role | Location",
                "AI Agent Intern | Shanghai",
            ],
            [fragment.text for fragment in parsed.fragments],
        )
        self.assertEqual(
            [1, 1, 1, 2, 2, 2],
            [
                fragment.location.slide_number
                for fragment in parsed.fragments
            ],
        )
        self.assertEqual(
            ("Open Positions",),
            parsed.fragments[5].location.heading_path,
        )
        self.assertEqual(1, parsed.fragments[5].location.table_number)
        self.assertEqual(
            2,
            parsed.fragments[5].location.table_row_number,
        )

    def test_registry_dispatches_pptx_parser(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "notice.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )
            slide.shapes.title.text = "Application Process"
            presentation.save(path)
            artifact = inspect_document_artifact(path, "source-pptx")

            parsed = DocumentParserRegistry(
                [PptxDocumentParser()]
            ).parse(artifact)

        self.assertEqual(
            "Application Process",
            parsed.fragments[0].text,
        )

    def test_empty_presentation_has_no_extractable_content(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "empty.pptx"
            Presentation().save(path)
            artifact = inspect_document_artifact(path, "source-empty")
            service = DocumentParsingService(
                DocumentParserRegistry([PptxDocumentParser()])
            )

            result = service.parse(artifact)

        self.assertEqual(
            "no-extractable-content",
            result.failure.code,
        )
        self.assertEqual(
            "use-fallback-extractor-or-manual-review",
            result.failure.operator_action,
        )

    def test_corrupt_pptx_has_stable_failure_classification(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "corrupt.pptx"
            path.write_bytes(b"not-an-office-presentation")
            artifact = inspect_document_artifact(path, "source-corrupt")
            service = DocumentParsingService(
                DocumentParserRegistry([PptxDocumentParser()])
            )

            result = service.parse(artifact)

        self.assertEqual(
            "invalid-presentation",
            result.failure.code,
        )

    def test_rejects_non_pptx_artifact_when_called_directly(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "notice.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )
            presentation.save(path)
            pptx_artifact = inspect_document_artifact(
                path,
                "source-pptx",
            )
            xlsx_artifact = replace(
                pptx_artifact,
                document_format="xlsx",
            )

            with self.assertRaisesRegex(
                DocumentParsingError,
                "only accepts pptx",
            ):
                PptxDocumentParser().parse(xlsx_artifact)


if __name__ == "__main__":
    unittest.main()
