import unittest
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app.application.documents.document_parsing import (
    DocumentParserRegistry,
    DocumentParsingService,
)
from app.domain.documents.document_ingestion import (
    DocumentParsingError,
    inspect_document_artifact,
)
from app.infrastructure.documents.docx_document_parser import (
    DocxDocumentParser,
)


class DocxDocumentParserTest(unittest.TestCase):
    def test_extracts_paragraphs_headings_and_table_rows(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "recruitment.docx"
            document = Document()
            document.add_heading("2027 Campus Recruitment", level=1)
            document.add_paragraph("Open to graduating students.")
            document.add_heading("AI Agent Intern", level=2)
            document.add_paragraph("Experience with RAG is preferred.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Location"
            table.cell(0, 1).text = "Degree"
            table.cell(1, 0).text = "Shanghai"
            table.cell(1, 1).text = "Bachelor"
            document.save(path)
            artifact = inspect_document_artifact(
                path,
                "https://careers.example/recruitment.docx",
            )

            parsed = DocxDocumentParser().parse(artifact)

        self.assertEqual(
            [
                "2027 Campus Recruitment",
                "Open to graduating students.",
                "AI Agent Intern",
                "Experience with RAG is preferred.",
                "Location | Degree",
                "Shanghai | Bachelor",
            ],
            [fragment.text for fragment in parsed.fragments],
        )
        self.assertEqual(
            ("2027 Campus Recruitment", "AI Agent Intern"),
            parsed.fragments[3].location.heading_path,
        )
        self.assertEqual(1, parsed.fragments[5].location.table_number)
        self.assertEqual(
            2,
            parsed.fragments[5].location.table_row_number,
        )

    def test_registry_dispatches_docx_parser(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "notice.docx"
            document = Document()
            document.add_paragraph("Application deadline: Friday.")
            document.save(path)
            artifact = inspect_document_artifact(path, "source-docx")
            registry = DocumentParserRegistry(
                [DocxDocumentParser()]
            )

            parsed = registry.parse(artifact)

        self.assertEqual(
            "Application deadline: Friday.",
            parsed.fragments[0].text,
        )

    def test_corrupt_docx_has_stable_failure_classification(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "corrupt.docx"
            path.write_bytes(b"not-an-office-package")
            artifact = inspect_document_artifact(path, "source-corrupt")
            service = DocumentParsingService(
                DocumentParserRegistry([DocxDocumentParser()])
            )

            result = service.parse(artifact)

        self.assertFalse(result.succeeded)
        self.assertEqual(
            "invalid-document-package",
            result.failure.code,
        )
        self.assertEqual(
            "redownload-or-manual-review",
            result.failure.operator_action,
        )

    def test_rejects_non_docx_artifact_when_called_directly(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "notice.docx"
            document = Document()
            document.add_paragraph("Content")
            document.save(path)
            docx_artifact = inspect_document_artifact(
                path,
                "source-docx",
            )
            pdf_artifact = replace(
                docx_artifact,
                document_format="pdf",
            )

            with self.assertRaisesRegex(
                DocumentParsingError,
                "only accepts docx",
            ):
                DocxDocumentParser().parse(pdf_artifact)


if __name__ == "__main__":
    unittest.main()
